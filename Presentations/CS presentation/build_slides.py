"""Build CS_presentation_IFTY.pptx (16:9, formal academic style, Cambria).

Run with the project venv:
    venv/bin/python "Presentations/CS presentation/build_slides.py"
Each slide is one function; slides are added in the order of SLIDES.
Preview of one slide:  build_slides.py N out.pptx
The speaker script is written to CS_presentation_IFTY_script.md (it is not stored in the deck).
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
import sys
sys.path.insert(0, str(HERE))  # make_figures.py lives next to this file
OUT = HERE / "CS_presentation_IFTY.pptx"
SCRIPT_OUT = HERE / "CS_presentation_IFTY_script.md"

# Palette: deep navy, muted gold, warm ivory panels, dark ink text
NAVY = RGBColor(0x0B, 0x25, 0x45)
GOLD = RGBColor(0xB0, 0x8D, 0x57)
IVORY = RGBColor(0xF6, 0xF3, 0xEC)
INK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80)
PALE = RGBColor(0xC9, 0xD3, 0xE0)  # light text on navy
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WINE = RGBColor(0x8B, 0x2E, 0x2E)  # muted red for "does not improve"
SOFT_NAVY = RGBColor(0x2E, 0x4A, 0x70)  # middle step between ivory and navy
LIGHT_GOLD = RGBColor(0xD9, 0xC4, 0x9F)
FONT = "Cambria"
FONT_JA = "Yu Mincho"
CJK = re.compile(r"[\u3000-\u30ff\u4e00-\u9fff\uff00-\uffef]")

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.75)
NUM = 0  # number of the slide being built (set in main)
SHORT_TITLE = "An Empirical Analysis of Entangled Enhancement in an Autonomous Driving Pipeline"


# ───────────────────────── helpers ─────────────────────────

BOLD = re.compile(r"(\*\*.+?\*\*)")
SUB = re.compile(r"(_\{[^}]*\})")


def add_runs(p, line, size, color, bold, italic, spacing=None):
    """Mini markup inside a line: **bold** and _{subscript} (subscripts may sit inside bold)."""
    for seg in BOLD.split(line):
        if not seg:
            continue
        b = bold
        if seg.startswith("**") and seg.endswith("**"):
            seg, b = seg[2:-2], True
        for part in SUB.split(seg):
            if not part:
                continue
            sub = part.startswith("_{")
            if sub:
                part = part[2:-1]
            r = p.add_run()
            r.text = part
            r.font.name, r.font.size, r.font.bold, r.font.italic = FONT, Pt(size), b, italic
            r.font.color.rgb = color
            if CJK.search(part):  # Cambria has no Japanese glyphs; Yu Mincho ships with Office
                latin = r.font._element.find(qn("a:latin"))
                latin.addnext(latin.makeelement(qn("a:ea"), {"typeface": FONT_JA}))
            if sub:
                r.font._element.set("baseline", "-25000")
            if spacing is not None:
                r.font._element.set("spc", str(int(spacing * 100)))


def text(slide, x, y, w, h, s, size, color=INK, bold=False, italic=False, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, spacing=None, line_spacing=None):
    """One text box with one paragraph per line of s. spacing = letter spacing in pt."""
    tf = slide.shapes.add_textbox(*map(int, (x, y, w, h))).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        add_runs(p, line, size, color, bold, italic, spacing)
    return tf


def arrow(slide, x1, y1, x2, y2, color=NAVY, width=1.75, dashed=False):
    """Straight connector with a triangular arrowhead at (x2, y2)."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *map(int, (x1, y1, x2, y2)))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def model_box(slide, x, y, w, h, name, role, fill=NAVY, name_color=WHITE, role_color=PALE, line=None):
    """Box for one ML model: symbol on top, role underneath."""
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, line=line, line_w=1.5)
    text(slide, x, y + Inches(0.1), w, h * 0.5, name, 24, name_color, bold=True)
    text(slide, x, y + h * 0.52, w, h * 0.4, role, 13, role_color, italic=True)


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=0.75):
    s = slide.shapes.add_shape(kind, *map(int, (x, y, w, h)))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def bar(slide, x, y, w, h, color=NAVY):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=color)


def ornament(slide, cx, y, half_w, color=GOLD):
    """Thin rule with a small diamond in the middle."""
    d = Inches(0.13)
    bar(slide, cx - half_w, y + d / 2 - Pt(0.5), half_w - d, Pt(1), color)
    bar(slide, cx + d, y + d / 2 - Pt(0.5), half_w - d, Pt(1), color)
    shape(slide, MSO_SHAPE.DIAMOND, cx - d / 2, y, d, d, fill=color)


SCRIPT = []          # (slide number, title, speaker script); written to SCRIPT_OUT, not into the deck
TITLE = ""           # title of the slide being built (set by content_slide, cover and thanks)


def notes(slide, s):
    """Collect the speaker script of a slide for the separate Markdown file."""
    SCRIPT.append((NUM, TITLE, s))


def content_slide(prs, title, section=""):
    """Common frame: navy title band, gold rule, section label, footer with slide number."""
    global TITLE
    TITLE = title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    band = Inches(1.15)
    bar(s, 0, 0, W, band, NAVY)
    bar(s, 0, band, W, Pt(3), GOLD)
    text(s, MARGIN, 0, Inches(9.5), band, title, 30, WHITE, bold=True, align=PP_ALIGN.LEFT)
    if section:
        text(s, W - MARGIN - Inches(3.5), 0, Inches(3.5), band, section.upper(), 11, GOLD,
             align=PP_ALIGN.RIGHT, spacing=2)
    # footer
    fy = H - Inches(0.55)
    bar(s, MARGIN, fy, W - 2 * MARGIN, Pt(0.75), GOLD)
    text(s, MARGIN, fy + Inches(0.1), Inches(9), Inches(0.35),
         SHORT_TITLE, 11, GREY, italic=True, align=PP_ALIGN.LEFT)
    text(s, W - MARGIN - Inches(1), fy + Inches(0.1), Inches(1), Inches(0.35), str(NUM), 12, NAVY,
         bold=True, align=PP_ALIGN.RIGHT)
    return s


# ───────────────────────── slides ─────────────────────────

def cover(prs):
    global TITLE
    TITLE = "Cover"
    s = prs.slides.add_slide(prs.slide_layouts[6])
    I = Inches
    bar(s, 0, 0, W, H, NAVY)
    inset = I(0.35)
    shape(s, MSO_SHAPE.RECTANGLE, inset, inset, W - 2 * inset, H - 2 * inset, line=GOLD, line_w=1)
    shape(s, MSO_SHAPE.RECTANGLE, inset + I(0.08), inset + I(0.08),
          W - 2 * inset - I(0.16), H - 2 * inset - I(0.16), line=GOLD, line_w=0.5)
    text(s, I(1), I(0.85), W - I(2), I(0.4), "0ALD515  ·  RESEARCH IN COMPUTER SCIENCE B", 13, GOLD, spacing=3)
    text(s, I(0.7), I(1.3), W - I(1.4), I(1.65),
         "An Empirical Analysis of Entangled Enhancement\nin an Autonomous Driving Pipeline", 34, WHITE, bold=True)
    ornament(s, W / 2, I(3.1), I(2.6))

    # presenter (left, aligned toward the centre line) and supervisor (right)
    cx, gap, cw = W / 2, I(0.45), I(5.0)
    lx, rx = cx - gap - cw, cx + gap
    bar(s, cx - Pt(0.5), I(3.65), Pt(1), I(2.2), GOLD)
    R, L = PP_ALIGN.RIGHT, PP_ALIGN.LEFT
    text(s, lx, I(3.6), cw, I(0.3), "PRESENTER", 11, GOLD, bold=True, align=R, spacing=2)
    text(s, lx, I(3.92), cw, I(0.5), "Rashedul Arefin Ifty", 24, WHITE, align=R)
    text(s, lx, I(4.45), cw, I(0.35), "Student ID: 202620721", 14, PALE, align=R)
    text(s, lx, I(4.85), cw, I(1.0), "理工情報生命学術院\nシステム情報工学研究群\n情報理工学位プログラム", 13, PALE,
         align=R, anchor=MSO_ANCHOR.TOP)
    text(s, rx, I(3.6), cw, I(0.3), "SUPERVISOR", 11, GOLD, bold=True, align=L, spacing=2)
    text(s, rx, I(3.92), cw, I(0.5), "Fumio Machida", 24, WHITE, align=L)
    text(s, rx, I(4.45), cw, I(0.35), "Associate Professor", 14, PALE, align=L)
    text(s, rx, I(4.85), cw, I(1.0),
         "Department of Systems and Information Engineering\nDependable AI Systems Laboratory", 13, PALE,
         align=L, anchor=MSO_ANCHOR.TOP)

    text(s, I(1), I(6.25), W - I(2), I(0.4), "University of Tsukuba", 16, GOLD, spacing=1.5)
    notes(s, "My name is Rashedul Arefin Ifty, from the Dependable AI Systems Laboratory. My supervisor is "
             "Associate Professor Fumio Machida. In this presentation, I report an empirical analysis of "
             "entangled enhancement in an autonomous driving pipeline.")


def outline(prs):
    s = content_slide(prs, "Outline")
    items = [("Background and Existing Studies", "Entangled enhancement and the limitation of existing studies"),
             ("Objective and Research Questions", "What this study investigates, and why"),
             ("Target Pipeline and Experimental Setup", "Models, data, updates and evaluation procedure"),
             ("Results", "One answer for each research question"),
             ("Conclusion and Future Work", "Summary and the next step toward early detection")]
    numerals = ["I", "II", "III", "IV", "V"]
    left, top, step, ch = Inches(1.6), Inches(1.6), Inches(1.0), Inches(0.82)
    cw = W - 2 * left
    for i, (head, sub) in enumerate(items):
        y = top + i * step
        shape(s, MSO_SHAPE.RECTANGLE, left, y, cw, ch, fill=IVORY)
        bar(s, left, y, Inches(1.0), ch, NAVY)
        text(s, left, y, Inches(1.0), ch, numerals[i], 24, GOLD, bold=True)
        text(s, left + Inches(1.35), y + Inches(0.08), cw - Inches(1.6), Inches(0.42), head, 22, NAVY, bold=True,
             align=PP_ALIGN.LEFT)
        text(s, left + Inches(1.35), y + Inches(0.47), cw - Inches(1.6), Inches(0.3), sub, 14, GREY, italic=True,
             align=PP_ALIGN.LEFT)
    notes(s, 'The presentation has five parts. They are the background, the objective, the setup, the results, and'
             ' the conclusion.')


def background(prs):
    s = content_slide(prs, "Background: Entangled Enhancement", "I · Background")
    I = Inches

    # Left: the MLS as a pipeline, and the retraining of C1
    text(s, MARGIN, I(1.45), I(7), I(0.35), "MACHINE LEARNING SYSTEM (MLS) AS A PIPELINE", 12, GOLD,
         bold=True, align=PP_ALIGN.LEFT, spacing=1.5)
    text(s, MARGIN, I(1.78), I(7.2), I(0.3),
         "Each model uses the output of the upstream model before it. The models after it are downstream.",
         12, GREY, italic=True, align=PP_ALIGN.LEFT)
    by, bw, bh = I(2.3), I(1.65), I(0.95)
    xs = [I(1.25), I(3.35), I(5.45)]
    for x, (nm, role) in zip(xs, [("C_{1}", "Perception model"), ("C_{2}", "Prediction model"),
                                  ("C_{3}", "Planning model")]):
        model_box(s, x, by, bw, bh, nm, role)
    mid = by + bh / 2
    arrow(s, MARGIN, mid, xs[0], mid)
    arrow(s, xs[0] + bw, mid, xs[1], mid)
    arrow(s, xs[1] + bw, mid, xs[2], mid)
    arrow(s, xs[2] + bw, mid, xs[2] + bw + I(0.55), mid)
    text(s, MARGIN - I(0.1), mid - I(0.4), I(0.6), I(0.3), "image", 12, GREY, italic=True)
    text(s, xs[2] + bw, mid - I(0.4), I(0.6), I(0.3), "path", 12, GREY, italic=True)

    # retraining C1 -> C1'
    ry = I(3.9)
    arrow(s, xs[0] + bw / 2, by + bh, xs[0] + bw / 2, ry, GOLD, 2, dashed=True)
    text(s, xs[0] + bw / 2 + I(0.12), by + bh + I(0.12), I(1.3), I(0.5), "retraining", 13, GOLD, italic=True,
         align=PP_ALIGN.LEFT)
    model_box(s, xs[0], ry, bw, bh, "C_{1}′", "Updated model", fill=IVORY, name_color=NAVY,
              role_color=GREY, line=GOLD)
    # outcome of the retraining
    ox = I(3.15)
    shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, ox, ry + I(0.2), I(0.24), I(0.2), fill=GOLD)
    text(s, ox + I(0.4), ry + I(0.08), I(4.4), I(0.45), "The performance of C_{1}′ improves.", 15,
         INK, align=PP_ALIGN.LEFT)
    tri = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, ox, ry + I(0.72), I(0.24), I(0.2), fill=WINE)
    tri.rotation = 180
    text(s, ox + I(0.4), ry + I(0.6), I(4.4), I(0.45), "The performance of the MLS does not improve.", 15,
         WINE, align=PP_ALIGN.LEFT)

    # Right: definition
    dx, dy, dw, dh = I(8.15), I(1.5), W - MARGIN - I(8.15), I(3.4)
    shape(s, MSO_SHAPE.RECTANGLE, dx, dy, dw, dh, fill=IVORY)
    bar(s, dx, dy, Pt(5), dh, GOLD)
    text(s, dx + I(0.35), dy + I(0.2), dw - I(0.6), I(0.35), "DEFINITION", 12, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=2)
    text(s, dx + I(0.35), dy + I(0.65), dw - I(0.6), I(2.0),
         "**Entangled enhancement** is a retraining of one ML model that improves the performance "
         "of that model but fails to improve the performance of the whole MLS.",
         19, INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1)
    text(s, dx + I(0.35), dy + dh - I(0.55), dw - I(0.6), I(0.35),
         "Wang and Machida, COMPSAC 2024 and ISSRE 2025", 12, GREY, italic=True, align=PP_ALIGN.LEFT)

    # Bottom: the two causes
    cy, chh = I(5.15), I(1.6)
    cw = (W - 2 * MARGIN - I(0.35)) / 2
    causes = [("First cause: adaptation of downstream models",
               "A downstream model is trained on the output of the original upstream model, including its errors. "
               "The downstream model learns to work with these errors. The retraining changes these errors."),
              ("Second cause: uneven improvement",
               "The updated model improves on average. The updated model can still become worse on the inputs "
               "that matter to the downstream models.")]
    for k, (head, body) in enumerate(causes):
        cx = MARGIN + k * (cw + I(0.35))
        shape(s, MSO_SHAPE.RECTANGLE, cx, cy, cw, chh, fill=WHITE, line=NAVY, line_w=1)
        bar(s, cx, cy, cw, I(0.48), NAVY)
        text(s, cx + I(0.25), cy, cw - I(0.5), I(0.48), head, 15, WHITE, bold=True, align=PP_ALIGN.LEFT)
        text(s, cx + I(0.25), cy + I(0.58), cw - I(0.5), chh - I(0.68), body, 14, INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

    notes(s, 'A machine learning system, or MLS, connects several ML models sequentially. Each model uses the '
             'output of the upstream model before it. In our pipeline, the perception model C1 finds the objects, '
             'the prediction model C2 predicts their paths, and the planning model C3 plans the path of the '
             'vehicle. Retraining C1 gives the updated model C1 prime. Wang and Machida define entangled '
             'enhancement as a retraining that improves one ML model but fails to improve the whole MLS. Two '
             'causes can produce it. In the first cause, a downstream model is trained on the output of the '
             'original upstream model, including its errors. The downstream model learns to work with these '
             'errors. The retraining changes these errors. In the second cause, the updated model improves on '
             'average but becomes worse on the inputs that matter downstream.')


def table(slide, x, y, col_w, rows, row_h, header_h, size=14):
    """Formal table: navy header row, alternating ivory and white body rows.
    rows[0] is the header; cells use the same mini markup as text()."""
    shp = slide.shapes.add_table(len(rows), len(col_w), int(x), int(y), int(sum(col_w)),
                                 int(header_h + row_h * (len(rows) - 1)))
    tbl = shp.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = int(cw)
    tbl.rows[0].height = int(header_h)
    for i in range(1, len(rows)):
        tbl.rows[i].height = int(row_h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (IVORY if i % 2 else WHITE)
            c.margin_left = c.margin_right = Inches(0.15)
            c.margin_top = c.margin_bottom = Inches(0.06)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            for k, line in enumerate(val.split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT
                if i == 0:
                    add_runs(p, line, size, GOLD, True, False, spacing=1)
                elif k > 0:  # second line of a cell: small italic detail (venue, year)
                    add_runs(p, line, size - 2, GREY, False, True)
                else:
                    add_runs(p, line, size, NAVY if j == 0 else INK, j == 0, False)
    return tbl


def existing_studies(prs):
    s = content_slide(prs, "Existing Studies and Their Limitation", "I · Background")
    I = Inches
    rows = [["STUDY", "WHAT THE STUDY SHOWS", "WHAT REMAINS OPEN"],
            ["Sculley et al.\nNIPS 2015",
             "A change to one model of an MLS changes the behaviour of every model that uses its output.",
             "The study describes the risk. The study does not measure how often the risk occurs."],
            ["Wu et al.\nNeurIPS 2021",
             "Improving one model of a system can make the whole system worse.",
             "The study does not examine a pipeline of perception, prediction and planning models."],
            ["Wang and Machida\nCOMPSAC 2024, ISSRE 2025",
             "Models of when an MLS should retrain each of its ML models.",
             "The models take the probability of entangled enhancement as an input. "
             "Nobody has measured this probability in a driving pipeline."],
            ["Xu et al.\nICRA 2024",
             "The prediction error increases when the input comes from a real perception model "
             "instead of the ground truth, which is the objects labelled by hand.",
             "The study does not examine a retraining of the perception model."]]
    tw = W - 2 * MARGIN
    table(s, MARGIN, I(1.5), [I(2.9), (tw - I(2.9)) / 2, (tw - I(2.9)) / 2], rows, I(0.78), I(0.5))

    # research gap
    gy, gh = I(5.35), I(1.2)
    bar(s, MARGIN, gy, tw, gh, NAVY)
    bar(s, MARGIN, gy, Pt(6), gh, GOLD)
    text(s, MARGIN + I(0.4), gy, I(2.0), gh, "RESEARCH\nGAP", 15, GOLD, bold=True, align=PP_ALIGN.LEFT,
         spacing=2)
    text(s, MARGIN + I(2.5), gy, tw - I(2.9), gh,
         "No existing study measures how often entangled enhancement occurs in an autonomous driving "
         "pipeline of ML models, or which of the two causes produces it.",
         17, WHITE, align=PP_ALIGN.LEFT)

    notes(s, 'Sculley et al. describe how a change to one model affects every later model, but not how often this '
             'effect occurs. Wu et al. show that improving one model can make the whole system worse, but not in a'
             ' driving pipeline. Wang and Machida model when to retrain each model. Their models need the '
             'probability of entangled enhancement, which nobody has measured in a driving pipeline. Xu et al. '
             'show that the prediction error increases when the input comes from a real perception model instead '
             'of the ground truth, which is the objects labelled by hand. They do not study a retraining. No study'
             ' therefore measures how often entangled enhancement occurs in a driving pipeline, or which cause '
             'produces it.')


def objective(prs):
    s = content_slide(prs, "Objective and Research Questions", "II · Objective")
    I = Inches
    tw = W - 2 * MARGIN

    # objective banner
    oy, oh = I(1.45), I(1.0)
    shape(s, MSO_SHAPE.RECTANGLE, MARGIN, oy, tw, oh, fill=IVORY)
    bar(s, MARGIN, oy, Pt(6), oh, GOLD)
    text(s, MARGIN + I(0.35), oy, I(1.7), oh, "OBJECTIVE", 14, GOLD, bold=True, align=PP_ALIGN.LEFT, spacing=2)
    text(s, MARGIN + I(2.1), oy, tw - I(2.4), oh,
         "In this study, we empirically investigate whether a retraining that improves the perception model "
         "C_{1} increases the error of the prediction model C_{2} and the planning model C_{3}.",
         17, INK, align=PP_ALIGN.LEFT)

    # three research questions
    cards = [("RQ1 · OCCURRENCE",
              "Does a retraining that improves C_{1} increase the error of C_{2} or C_{3}?",
              "The models of Wang and Machida need the probability of entangled enhancement. "
              "Nobody has measured this probability in a driving pipeline.",
              "H1", "A retraining that improves C_{1} can increase the prediction error or the planning error."),
             ("RQ2 · CAUSE",
              "Does the increase come from the adaptation of C_{2} and C_{3} to the errors of C_{1}?",
              "Only training on the output of C_{1} lets C_{2} and C_{3} learn the errors of C_{1}. "
              "If this adaptation is the cause, retraining C_{2} and C_{3} again can remove the increase.",
              "H2", "The increase is larger when C_{2} and C_{3} are trained on the output of C_{1} "
                    "than on the ground truth."),
             ("RQ3 · DETECTION",
              "Can we detect entangled enhancement without ground truth?",
              "After deployment, new data have no ground truth. "
              "A larger change in the input of a model can cause a larger increase of its error.",
              "H3", "Retrainings with entangled enhancement change the outputs of C_{1} and C_{2} "
                    "more than other retrainings.")]
    gap = I(0.3)
    cw = (tw - 2 * gap) / 3
    cy, ch = I(2.7), I(4.05)
    for k, (head, q, why, hn, hyp) in enumerate(cards):
        cx = MARGIN + k * (cw + gap)
        shape(s, MSO_SHAPE.RECTANGLE, cx, cy, cw, ch, fill=WHITE, line=NAVY, line_w=1)
        bar(s, cx, cy, cw, I(0.48), NAVY)
        text(s, cx + I(0.2), cy, cw - I(0.4), I(0.48), head, 13, GOLD, bold=True, align=PP_ALIGN.LEFT, spacing=1.5)
        text(s, cx + I(0.2), cy + I(0.6), cw - I(0.4), I(0.85), q, 16, NAVY, bold=True, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        text(s, cx + I(0.2), cy + I(1.5), cw - I(0.4), I(0.3), "WHY WE ASK", 10, GOLD, bold=True,
             align=PP_ALIGN.LEFT, spacing=1.5)
        text(s, cx + I(0.2), cy + I(1.78), cw - I(0.4), I(1.1), why, 13, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        hy = cy + ch - I(1.12)
        shape(s, MSO_SHAPE.RECTANGLE, cx + I(0.12), hy, cw - I(0.24), I(1.0), fill=IVORY)
        bar(s, cx + I(0.12), hy, Pt(4), I(1.0), GOLD)
        text(s, cx + I(0.3), hy, I(0.5), I(1.0), hn, 18, GOLD, bold=True, align=PP_ALIGN.LEFT)
        text(s, cx + I(0.8), hy + I(0.05), cw - I(1.05), I(0.9), hyp, 13, INK, italic=True, align=PP_ALIGN.LEFT)

    notes(s, 'In this study, we empirically investigate whether a retraining that improves the perception model C1'
             ' increases the error of the prediction model C2 and the planning model C3. We ask three research '
             'questions. RQ1 asks whether the error increases, because the models of Wang and Machida need this '
             'probability. H1 states that the error can increase. RQ2 asks whether the increase comes from the '
             'adaptation of C2 and C3 to the errors of C1. H2 states that the increase is larger after training on'
             ' the output of C1 than after training on the ground truth. RQ3 asks whether we can detect entangled '
             'enhancement without ground truth, because new data have no ground truth after deployment. H3 states '
             'that retrainings with entangled enhancement change the outputs of C1 and C2 more than other '
             'retrainings.')


def pipeline(prs):
    s = content_slide(prs, "Target Pipeline and Notation", "III · Setup")
    I = Inches

    # pipeline: input box and three model cards
    text(s, MARGIN, I(1.4), I(8), I(0.35), "AUTONOMOUS DRIVING PIPELINE OF THREE ML MODELS", 12, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.5)
    py, ph = I(1.85), I(2.35)
    shape(s, MSO_SHAPE.RECTANGLE, MARGIN, py + I(0.6), I(1.3), I(1.15), fill=IVORY, line=GOLD, line_w=1)
    text(s, MARGIN, py + I(0.6), I(1.3), I(1.15), "Front-camera\nimage", 14, NAVY, bold=True)
    cards = [("C_{1} · Perception model", "YOLOv11n object detector",
              "**Output:** objects with class, position and confidence"),
             ("C_{2} · Prediction model", "AutoBot-Ego style transformer",
              "**Input:** objects in the last four labelled images (1.5 s)\n**Output:** six possible paths of 3 s for "
              "each object"),
             ("C_{3} · Planning model", "PlanT style transformer",
              "**Input:** objects with their most likely path, driving command, path of the ego vehicle over "
              "the last 1.5 s (ego path)\n**Output:** planned path of the ego vehicle for 3 s")]
    cw, gap = I(3.15), I(0.35)
    x0 = MARGIN + I(1.3) + gap
    for k, (head, model, io) in enumerate(cards):
        cx = x0 + k * (cw + gap)
        shape(s, MSO_SHAPE.RECTANGLE, cx, py, cw, ph, fill=WHITE, line=NAVY, line_w=1)
        bar(s, cx, py, cw, I(0.5), NAVY)
        text(s, cx + I(0.18), py, cw - I(0.36), I(0.5), head, 15, WHITE, bold=True, align=PP_ALIGN.LEFT)
        text(s, cx + I(0.18), py + I(0.6), cw - I(0.36), I(0.35), model, 14, NAVY, italic=True,
             align=PP_ALIGN.LEFT)
        bar(s, cx + I(0.18), py + I(1.0), I(0.6), Pt(1.5), GOLD)
        text(s, cx + I(0.18), py + I(1.12), cw - I(0.36), ph - I(1.2), io, 12.5, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        prev_right = MARGIN + I(1.3) if k == 0 else cx - gap
        arrow(s, prev_right, py + ph / 2 + I(0.05), cx, py + ph / 2 + I(0.05))

    # notation
    ny = I(4.5)
    text(s, MARGIN, ny - I(0.15), I(8), I(0.35), "NOTATION FOR ONE RETRAINING OF C_{1}", 12, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.5)
    sy, sh_ = ny + I(0.3), I(1.95)
    syms = [("δ_{1}", "Gain of C_{1} in mAP50, the standard accuracy measure of object detection"),
            ("Δ_{2}", "Increase of the prediction error minADE: distance between the predicted and the true path, "
                      "for the best of the six paths"),
            ("Δ_{3}", "Increase of the planning error: distance between the planned position and the position "
                      "of the human driver at 3 s")]
    mw, mgap = I(2.65), I(0.25)
    for k, (sym, desc) in enumerate(syms):
        mx = MARGIN + k * (mw + mgap)
        shape(s, MSO_SHAPE.RECTANGLE, mx, sy, mw, sh_, fill=IVORY)
        text(s, mx + I(0.2), sy + I(0.1), I(1.2), I(0.6), sym, 28, GOLD, bold=True, align=PP_ALIGN.LEFT)
        text(s, mx + I(0.2), sy + I(0.72), mw - I(0.35), sh_ - I(0.8), desc, 12.5, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
    # condition
    qx = MARGIN + 3 * (mw + mgap)
    qw = W - MARGIN - qx
    bar(s, qx, sy, qw, sh_, NAVY)
    bar(s, qx, sy, qw, Pt(4), GOLD)
    text(s, qx + I(0.25), sy + I(0.15), qw - I(0.5), I(0.35), "ENTANGLED ENHANCEMENT", 11, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.5)
    text(s, qx + I(0.25), sy + I(0.5), qw - I(0.5), I(0.5), "A retraining shows entangled enhancement when",
         13, PALE, italic=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    text(s, qx + I(0.1), sy + I(0.95), qw - I(0.2), I(0.9), "δ_{1} > 0   and\n(Δ_{2} > 0   or   Δ_{3} > 0)",
         19, WHITE, bold=True)

    notes(s, 'We use YOLOv11n as the perception model C1. C1 finds the objects in the front-camera image. We '
             'employ an AutoBot-Ego style transformer as the prediction model C2. C2 reads the objects of the last'
             ' 1.5 seconds and outputs six possible paths of three seconds for each object. For the planning model'
             ' C3, we employ a PlanT style transformer. C3 receives the objects with their most likely path, the '
             'driving command, and the ego path. The ego path is the path of the ego vehicle, the vehicle that '
             'carries the camera, over the last 1.5 seconds. For each retraining, delta one is the gain of C1 in '
             'accuracy. Delta two is the increase of the prediction error. Delta three is the increase of the '
             'planning error at three seconds.')


def step_box(slide, x, y, w, h, s, fill=WHITE, line=NAVY, size=14, color=INK):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, line=line, line_w=1)
    text(slide, x + Inches(0.12), y, w - Inches(0.24), h, s, size, color)


def training_modes(prs):
    s = content_slide(prs, "Two Training Modes of C_{2} and C_{3}", "III · Setup")
    I = Inches

    lanes = [("GT TRAINING", "C_{2} and C_{3} never see\nthe errors of C_{1}",
              "**Ground truth (GT)**\nobjects labelled by hand"),
             ("C_{1}-OUTPUT TRAINING", "C_{2} and C_{3} learn\nthe errors of C_{1}",
              "**Output of C_{1}**\nincluding its errors")]
    lx, lw = MARGIN, I(2.55)
    b1x, bw = lx + lw + I(0.35), I(2.3)
    b2x = b1x + bw + I(0.4)
    tx = b2x + bw + I(0.45)
    tw_ = W - MARGIN - tx
    ly0, lh, lgap = I(1.55), I(1.2), I(0.3)
    for k, (name, sub, src) in enumerate(lanes):
        y = ly0 + k * (lh + lgap)
        bar(s, lx, y, lw, lh, NAVY)
        bar(s, lx, y, Pt(5), lh, GOLD)
        text(s, lx + I(0.25), y + I(0.12), lw - I(0.35), I(0.4), name, 13, GOLD, bold=True, align=PP_ALIGN.LEFT,
             spacing=1)
        text(s, lx + I(0.25), y + I(0.5), lw - I(0.35), I(0.62), sub, 12.5, PALE, italic=True,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
        step_box(s, b1x, y, bw, lh, src, fill=IVORY, line=GOLD)
        step_box(s, b2x, y, bw, lh, "**Train C_{2} and C_{3}**\nthen keep them fixed")
        arrow(s, b1x + bw, y + lh / 2, b2x, y + lh / 2)
        arrow(s, b2x + bw, y + lh / 2, tx, y + lh / 2)

    # common test
    th = 2 * lh + lgap
    bar(s, tx, ly0, tw_, th, NAVY)
    bar(s, tx, ly0, tw_, Pt(4), GOLD)
    text(s, tx + I(0.25), ly0 + I(0.2), tw_ - I(0.5), I(0.35), "TEST AFTER EACH RETRAINING", 12, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.5)
    text(s, tx + I(0.25), ly0 + I(0.65), tw_ - I(0.5), th - I(0.8),
         "Both versions of C_{2} and C_{3} receive the output of the same C_{1}′ on the same test scenes.\n"
         "Only the training differs.",
         15, WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1)

    # how we read the comparison
    ry, rh = I(4.55), I(2.15)
    tw = W - 2 * MARGIN
    shape(s, MSO_SHAPE.RECTANGLE, MARGIN, ry, tw, rh, fill=IVORY)
    bar(s, MARGIN, ry, Pt(6), rh, GOLD)
    text(s, MARGIN + I(0.35), ry + I(0.15), I(6), I(0.35), "HOW WE READ THE COMPARISON", 12, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.5)
    rows = [("The error increases only after C_{1}-output training.",
             "**First cause:** adaptation of C_{2} and C_{3} to the errors of C_{1}"),
            ("The error increases after both training modes.",
             "**Second cause:** uneven improvement of C_{1}′")]
    for k, (obs, cause) in enumerate(rows):
        y = ry + I(0.6) + k * I(0.7)
        text(s, MARGIN + I(0.35), y, I(5.1), I(0.55), obs, 16, INK, align=PP_ALIGN.LEFT)
        text(s, MARGIN + I(5.45), y, I(0.6), I(0.55), "→", 22, GOLD, bold=True)
        text(s, MARGIN + I(6.15), y, tw - I(6.3), I(0.55), cause, 16, NAVY, align=PP_ALIGN.LEFT)

    notes(s, 'To answer RQ2, we train C2 and C3 in two modes. In GT training, where GT means ground truth, C2 and '
             'C3 learn from the objects labelled by hand. They never see the errors of C1. In C1-output training, '
             'C2 and C3 learn from the output of C1, including its errors. In both modes, we train C2 and C3 once '
             'and keep them fixed. After each retraining, both versions receive the output of the same C1 prime. '
             'Only the training differs. If the error increases only after C1-output training, the cause is the '
             'adaptation to the errors of C1. If the error increases after both modes, the increase points to the '
             'uneven improvement of C1 prime.')


def setup(prs):
    s = content_slide(prs, "Experimental Setup", "III · Setup")
    I = Inches

    # data split (left)
    lx, lw = MARGIN, I(6.85)
    text(s, lx, I(1.4), lw, I(0.35), "DATA SET", 12, GOLD, bold=True, align=PP_ALIGN.LEFT, spacing=1.5)
    text(s, lx, I(1.72), lw, I(0.45),
         "350 of the 850 scenes of nuScenes (driving clips of about 20 s), front camera only, 4.6 GB",
         13.5, INK, italic=True, align=PP_ALIGN.LEFT)
    tiles = [("BOSTON 1", NAVY, WHITE, "100 scenes", "Training of C_{1}", "Two original C_{1}, with random seed 0 and 1"),
             ("BOSTON 2", NAVY, WHITE, "100 scenes", "Training of C_{2} and C_{3}",
              "86 training and 14 validation scenes, separate from Boston 1"),
             ("SINGAPORE 1", GOLD, WHITE, "100 scenes", "Retraining of C_{1}",
              "A new city, as in an update after deployment"),
             ("SINGAPORE 2", GOLD, WHITE, "50 scenes", "Test",
              "1,720 test samples, one per labelled image with 3 s of future")]
    tw_, th_, g = (lw - I(0.25)) / 2, I(1.6), I(0.22)
    for k, (city, band, fg, n, role, why) in enumerate(tiles):
        x = lx + (k % 2) * (tw_ + I(0.25))
        y = I(2.3) + (k // 2) * (th_ + g)
        shape(s, MSO_SHAPE.RECTANGLE, x, y, tw_, th_, fill=IVORY)
        bar(s, x, y, tw_, I(0.4), band)
        text(s, x + I(0.2), y, tw_ - I(0.4), I(0.4), city, 12.5, fg, bold=True, align=PP_ALIGN.LEFT, spacing=1.5)
        text(s, x + I(0.2), y, tw_ - I(0.4), I(0.4), n, 12.5, fg, italic=True, align=PP_ALIGN.RIGHT)
        text(s, x + I(0.2), y + I(0.5), tw_ - I(0.4), I(0.4), role, 16, NAVY, bold=True, align=PP_ALIGN.LEFT)
        if why:
            text(s, x + I(0.2), y + I(0.92), tw_ - I(0.4), I(0.6), why, 12, GREY, italic=True,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

    # retrainings (right)
    rx = lx + lw + I(0.4)
    rw = W - MARGIN - rx
    ry, rh = I(1.4), I(4.1)
    shape(s, MSO_SHAPE.RECTANGLE, rx, ry, rw, rh, fill=WHITE, line=NAVY, line_w=1)
    bar(s, rx, ry, rw, I(0.45), NAVY)
    text(s, rx + I(0.2), ry, rw - I(0.4), I(0.45), "RETRAININGS OF C_{1}", 12.5, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.5)
    text(s, rx + I(0.2), ry + I(0.55), I(1.2), I(0.8), "30", 44, GOLD, bold=True, align=PP_ALIGN.LEFT)
    text(s, rx + I(1.35), ry + I(0.58), rw - I(1.5), I(0.8),
         "retrainings\n= 2 original C_{1} × 15 settings", 15, NAVY, bold=True, align=PP_ALIGN.LEFT)
    text(s, rx + I(0.2), ry + I(1.42), rw - I(0.4), I(0.7),
         "**Original C_{1}:** YOLOv11n, 10 epochs on Boston 1, random seed 0 or 1",
         13, INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    rows = [["DATA USED", "EPOCHS", "SETTINGS"],
            ["100 %, 50 %, 25 %", "10", "3 × 3 seeds = 9"],
            ["100 %", "5 or 20", "2 × 3 seeds = 6"]]
    table(s, rx + I(0.2), ry + I(2.15), [I(1.65), I(0.95), rw - I(3.0)], rows, I(0.5), I(0.42), size=12.5)
    text(s, rx + I(0.2), ry + I(3.55), rw - I(0.4), I(0.45),
         "A random seed is the starting number for the random choices in training.", 11.5, GREY, italic=True,
         align=PP_ALIGN.LEFT)

    # environment strip
    ey, eh = I(5.82), I(0.75)
    tw = W - 2 * MARGIN
    bar(s, MARGIN, ey, tw, eh, NAVY)
    bar(s, MARGIN, ey, Pt(6), eh, GOLD)
    text(s, MARGIN + I(0.35), ey, I(2.2), eh, "ENVIRONMENT", 12.5, GOLD, bold=True, align=PP_ALIGN.LEFT, spacing=1.5)
    text(s, MARGIN + I(2.5), ey, tw - I(2.7), eh,
         "Kaggle notebook  ·  NVIDIA Tesla T4 GPU (15 GB)  ·  Python 3.13  ·  PyTorch 2.11 (CUDA 12.8)  ·  "
         "Ultralytics 8.4", 13, WHITE, align=PP_ALIGN.LEFT)

    notes(s, 'We use 350 of the 850 scenes of nuScenes. We split the scenes into four parts by city. Boston 1 '
             'trains C1. Boston 2 trains C2 and C3, which therefore see the real errors of C1 on unseen images. '
             'Singapore 1 is used for the retraining of C1, as in an update after deployment in a new city. '
             'Singapore 2 is the test set with 1,720 test samples. We train two original C1 with the random seeds '
             '0 and 1. We retrain each original C1 with 15 settings of data amount, epochs and seed. This gives 30'
             ' retrainings.')


def num_badge(slide, x, y, d, label, fill=NAVY, color=GOLD, size=16):
    c = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill)
    text(slide, x, y, d, d, label, size, color, bold=True)
    return c


def evaluation(prs):
    s = content_slide(prs, "Evaluation Procedure", "III · Setup")
    I = Inches

    # left: procedure
    lx, lw = MARGIN, I(5.55)
    text(s, lx, I(1.4), lw, I(0.35), "PROCEDURE FOR EACH RETRAINING", 12, GOLD, bold=True, align=PP_ALIGN.LEFT,
         spacing=1.5)
    steps = ["Run the pipeline with the original C_{1} on the 50 test scenes.",
             "Run the same fixed C_{2} and C_{3} with the retrained C_{1}′ on the same scenes.",
             "Compute δ_{1}, Δ_{2} and Δ_{3} as the differences between the two runs.",
             "Repeat for both training modes and all 30 retrainings."]
    for k, st in enumerate(steps):
        y = I(1.95) + k * I(0.98)
        num_badge(s, lx, y, I(0.5), str(k + 1))
        text(s, lx + I(0.7), y - I(0.08), lw - I(0.7), I(0.75), st, 15, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        if k < len(steps) - 1:
            bar(s, lx + I(0.25) - Pt(0.75), y + I(0.55), Pt(1.5), I(0.38), LIGHT_GOLD)
    shape(s, MSO_SHAPE.RECTANGLE, lx, I(5.9), lw, I(0.7), fill=IVORY)
    text(s, lx + I(0.2), I(5.9), lw - I(0.4), I(0.7),
         "A single positive Δ can come from chance. Each level of evidence removes one source of chance.",
         13, NAVY, italic=True, align=PP_ALIGN.LEFT)

    # right: three levels of evidence
    rx = lx + lw + I(0.45)
    rw = W - MARGIN - rx
    text(s, rx, I(1.4), rw, I(0.35), "THREE LEVELS OF EVIDENCE FOR AN INCREASE", 12, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.5)
    levels = [("LEVEL 1", "Any increase", "Δ_{2} > 0 or Δ_{3} > 0", IVORY, NAVY, INK),
              ("LEVEL 2", "Larger than seed variation",
               "Seed variation is the change in error when we replace the original C_{1} by the other original "
               "C_{1}. The two differ only in the random seed and have almost the same accuracy.\n"
               "**Threshold:  Δ_{2} > 0.112 m,  Δ_{3} > 0.018 m**", SOFT_NAVY, WHITE, PALE),
              ("LEVEL 3", "Statistically significant",
               "We randomly resample the 50 test scenes 2,000 times. The middle 95 % of the resampled "
               "increases lies above zero.", NAVY, WHITE, PALE)]
    hs = [I(0.95), I(1.95), I(1.3)]
    y = I(1.85)
    for (lv, head, body, fill, hc, bc), h in zip(levels, hs):
        bar(s, rx, y, rw, h, fill)
        bar(s, rx, y, Pt(6), h, GOLD)
        text(s, rx + I(0.3), y + I(0.12), I(1.0), I(0.35), lv, 11, GOLD, bold=True, align=PP_ALIGN.LEFT, spacing=1.5)
        text(s, rx + I(1.3), y + I(0.08), rw - I(1.5), I(0.42), head, 17, hc, bold=True, align=PP_ALIGN.LEFT)
        text(s, rx + I(1.3), y + I(0.5), rw - I(1.5), h - I(0.55), body, 13, bc, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        y += h + I(0.12)
    text(s, rx, y, rw, I(0.35), "weaker  →  stronger evidence", 11.5, GREY, italic=True, align=PP_ALIGN.RIGHT)

    notes(s, 'For each retraining, we run the pipeline with the original C1 and with C1 prime on the same 50 test '
             'scenes, with C2 and C3 fixed. The differences give the three deltas. A single positive delta can '
             'come from chance. We therefore count the increases at three levels. Level one counts any increase. '
             'Level two counts increases larger than the seed variation. The two original C1 differ only in the '
             'random seed and have almost the same accuracy. Replacing one by the other still changes the error by'
             ' up to 0.112 metres for prediction and 0.018 metres for planning. A smaller increase can therefore '
             'come from randomness. Level three resamples the test scenes 2,000 times. The middle 95 percent of '
             'the resampled increases must lie above zero.')


def model_check(prs):
    s = content_slide(prs, "Check of C_{2} and C_{3}", "III · Setup")
    I = Inches
    tw = W - 2 * MARGIN

    # why the check is needed
    wy, wh = I(1.45), I(0.85)
    shape(s, MSO_SHAPE.RECTANGLE, MARGIN, wy, tw, wh, fill=IVORY)
    bar(s, MARGIN, wy, Pt(6), wh, GOLD)
    text(s, MARGIN + I(0.35), wy, I(1.9), wh, "WHY WE\nCHECK", 12.5, GOLD, bold=True, align=PP_ALIGN.LEFT, spacing=1.5)
    text(s, MARGIN + I(2.1), wy, tw - I(2.35), wh,
         "If C_{2} or C_{3} were no better than a simple rule, a change of their error would say nothing about "
         "the retraining of C_{1}. We therefore compare both models with simple baselines on the test scenes.",
         15, INK, align=PP_ALIGN.LEFT)

    rows = [["MODEL", "INPUT", "ERROR", "SIMPLE BASELINE", "RESULT"],
            ["C_{2}", "Ground-truth objects", "0.77 m",
             "4.48 m: objects stay still\n1.01 m: objects keep their speed", "**Better**"],
            ["C_{2}", "Output of C_{1}", "2.39 to 2.95 m", "5.39 to 5.85 m: objects stay still", "**Better**"],
            ["C_{3} without ego path", "Output of C_{2}", "9.3 to 12.5 m",
             "9.25 m: average Boston path for each driving command", "Not better, not used"],
            ["C_{3} with ego path\nused in this study", "Output of C_{2} and the ego path of the last 1.5 s", "3.70 to 4.16 m",
             "9.25 m: average Boston path for each driving command", "**Better**"]]
    table(s, MARGIN, I(2.45), [I(2.4), I(3.05), I(1.55), tw - I(8.6), I(1.6)], rows, I(0.6), I(0.42), size=13)
    text(s, MARGIN, I(5.36), tw, I(0.3),
         "Error of C_{2}: minADE.  Error of C_{3}: distance to the position of the human driver at 3 s.  "
         "All values on the 50 test scenes.", 11.5, GREY, italic=True, align=PP_ALIGN.LEFT)

    # object-use check
    oy, oh = I(5.78), I(0.98)
    bar(s, MARGIN, oy, tw, oh, NAVY)
    bar(s, MARGIN, oy, Pt(6), oh, GOLD)
    text(s, MARGIN + I(0.35), oy, I(1.9), oh, "OBJECT-USE\nCHECK", 12.5, GOLD, bold=True, align=PP_ALIGN.LEFT,
         spacing=1.5)
    text(s, MARGIN + I(2.1), oy, tw - I(2.35), oh,
         "An ego-path input can let a planning model ignore the objects. "
         "Removing the objects from the input of C_{3} moves the planned position at 3 s by 0.73 to 1.65 m. "
         "This shift is above the threshold of 0.5 m. C_{3} still uses the objects.",
         14, WHITE, align=PP_ALIGN.LEFT)

    notes(s, 'Before the results, we check C2 and C3. If they were no better than a simple rule, a change of their'
             ' error would say nothing about the retraining. C2 has an error of 0.77 metres on ground-truth '
             'objects, against 4.48 metres when every object stays still. C3 without the ego path is not better '
             'than the average Boston path. We therefore do not use this variant. C3 with the ego path has about 4'
             ' metres, against 9.25 metres. Removing the objects still moves the plan of C3 by 0.73 to 1.65 '
             'metres. C3 therefore uses the objects.')


def takeaway(slide, y, h, label, body, size=15):
    """Navy answer box across the slide, gold label on the left."""
    tw = W - 2 * MARGIN
    bar(slide, MARGIN, y, tw, h, NAVY)
    bar(slide, MARGIN, y, Pt(6), h, GOLD)
    text(slide, MARGIN + Inches(0.35), y, Inches(2.0), h, label, 12.5, GOLD, bold=True, align=PP_ALIGN.LEFT,
         spacing=1.5)
    text(slide, MARGIN + Inches(2.4), y, tw - Inches(2.65), h, body, size, WHITE, align=PP_ALIGN.LEFT)


def result_prediction(prs):
    s = content_slide(prs, "Result: Prediction Model C_{2}", "IV · Results · RQ1, RQ2")
    I = Inches

    # figure (left)
    text(s, MARGIN, I(1.35), I(7.45), I(0.4),
         "We test whether the prediction error increases, and whether the training mode matters.",
         13.5, NAVY, italic=True, align=PP_ALIGN.LEFT)
    s.shapes.add_picture(str(HERE / "fig_delta2.png"), int(MARGIN), int(I(1.8)), width=int(I(6.95)))

    # right panel
    rx = MARGIN + I(7.75)
    rw = W - MARGIN - rx
    shape(s, MSO_SHAPE.RECTANGLE, rx, I(1.45), rw, I(0.95), fill=IVORY)
    bar(s, rx, I(1.45), Pt(4), I(0.95), GOLD)
    text(s, rx + I(0.25), I(1.5), rw - I(0.4), I(0.4), "C_{1}′ is better in all 30 retrainings", 14.5, NAVY,
         bold=True, align=PP_ALIGN.LEFT)
    text(s, rx + I(0.25), I(1.9), rw - I(0.4), I(0.4), "δ_{1} = +0.048 to +0.118 in mAP50", 13.5, INK,
         align=PP_ALIGN.LEFT)

    text(s, rx, I(2.6), rw, I(0.35), "RETRAININGS WITH Δ_{2} > 0  (OF 30)", 11.5, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.2)
    rows = [["LEVEL", "C_{1}-OUTPUT", "GT"],
            ["1. Any increase", "**23**", "5"],
            ["2. Above seed variation", "**11**", "0"],
            ["3. Significant", "**3**", "0"]]
    table(s, rx, I(2.95), [rw - I(1.95), I(1.3), I(0.65)], rows, I(0.45), I(0.42), size=13)

    shape(s, MSO_SHAPE.RECTANGLE, rx, I(4.85), rw, I(0.75), fill=WHITE, line=NAVY, line_w=1)
    text(s, rx + I(0.2), I(4.85), rw - I(0.4), I(0.75),
         "Δ_{2} is larger after C_{1}-output training than after GT training in **30 of 30** retrainings.",
         13.5, INK, align=PP_ALIGN.LEFT)

    takeaway(s, I(5.82), I(0.9), "ANSWER TO\nRQ1 AND RQ2",
             "H1 and H2 hold for C_{2}. The increases above seed variation appear only after C_{1}-output "
             "training. This result points to the first cause, the adaptation of C_{2} to the errors of C_{1}.",
             15)

    notes(s, 'C1 prime is more accurate than C1 in all 30 retrainings. The figure shows delta two, the increase of'
             ' the prediction error, for each retraining. The navy dots show C1-output training, and the gold '
             'diamonds show GT training. After C1-output training, delta two is positive in 23 retrainings, above '
             'the seed variation in 11, and significant in 3. After GT training, the counts are 5, 0 and 0. In all'
             ' 30 retrainings, delta two is larger after C1-output training, although both versions of C2 receive '
             'the same input. All 11 clear increases come from the original C1 with seed 0, on the left half. As a'
             ' result, we confirm H1 and H2 for C2.')


def result_planning(prs):
    s = content_slide(prs, "Result: Planning Model C_{3}", "IV · Results · RQ1, RQ2")
    I = Inches

    text(s, MARGIN, I(1.35), I(7.45), I(0.4),
         "We test whether the increase reaches the end of the pipeline, the planned path of C_{3}.",
         13.5, NAVY, italic=True, align=PP_ALIGN.LEFT)
    s.shapes.add_picture(str(HERE / "fig_delta3.png"), int(MARGIN), int(I(1.8)), width=int(I(6.95)))

    rx = MARGIN + I(7.75)
    rw = W - MARGIN - rx
    text(s, rx, I(1.45), rw, I(0.35), "RETRAININGS WITH Δ_{3} > 0  (OF 30)", 11.5, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.2)
    rows = [["LEVEL", "C_{1}-OUTPUT", "GT"],
            ["1. Any increase", "18", "**29**"],
            ["2. Above seed variation", "13", "**29**"],
            ["3. Significant", "1", "**10**"]]
    table(s, rx, I(1.8), [rw - I(1.95), I(1.3), I(0.65)], rows, I(0.45), I(0.42), size=13)

    cards = [("The increase is small",
              "Largest Δ_{3}: 0.13 m, against a planning error of 3.70 to 4.16 m."),
             ("The objects have a small effect on C_{3}",
              "Removing all objects moves the plan of C_{3} by only 0.73 to 1.65 m.")]
    y = I(3.75)
    for head, body in cards:
        shape(s, MSO_SHAPE.RECTANGLE, rx, y, rw, I(0.95), fill=IVORY)
        bar(s, rx, y, Pt(4), I(0.95), GOLD)
        text(s, rx + I(0.25), y + I(0.08), rw - I(0.4), I(0.35), head, 14, NAVY, bold=True, align=PP_ALIGN.LEFT)
        text(s, rx + I(0.25), y + I(0.42), rw - I(0.4), I(0.5), body, 12.5, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        y += I(1.05)

    takeaway(s, I(5.95), I(0.85), "ANSWER FOR C_{3}",
             "The planning error increases, but by at most 0.13 m. The increase is more frequent after GT "
             "training. The first cause cannot explain this increase.", 14.5)

    notes(s, 'Next, we test whether the increase reaches the planning model C3. The figure shows delta three, the '
             'increase of the planning error. After C1-output training, delta three is positive in 18 retrainings,'
             ' above the seed variation in 13, and significant in 1. After GT training, the counts are 29, 29 and '
             '10. The increase is small, at most 0.13 metres against a planning error of about 4 metres. Most of '
             'the plan of C3 comes from the ego path. The increase is also more frequent after GT training. A GT-'
             'trained C3 cannot have adapted to the errors of C1. The first cause therefore cannot explain this '
             'increase.')


def result_retraining(prs):
    s = content_slide(prs, "Result: Retraining C_{2} and C_{3} on the Output of C_{1}′", "IV · Results · RQ2")
    I = Inches

    s.shapes.add_picture(str(HERE / "fig_retrain.png"), int(MARGIN), int(I(1.4)), width=int(I(8.0)))

    rx = MARGIN + I(8.3)
    rw = W - MARGIN - rx
    cards = [("WHAT WE DO",
              "We retrain C_{2} and C_{3} on the output of C_{1}′. We test one retraining of each original C_{1}."),
             ("CHANGE IN ERROR",
              "Prediction error: −0.36 m and −0.16 m\nPlanning error: −0.44 m and −0.30 m"),
             ("RESULT",
              "Both errors fall below their values before the retraining of C_{1}.")]
    y = I(1.45)
    for head, body in cards:
        h = I(1.2)
        shape(s, MSO_SHAPE.RECTANGLE, rx, y, rw, h, fill=IVORY)
        bar(s, rx, y, Pt(4), h, GOLD)
        text(s, rx + I(0.25), y + I(0.1), rw - I(0.4), I(0.3), head, 11, GOLD, bold=True, align=PP_ALIGN.LEFT,
             spacing=1.5)
        text(s, rx + I(0.25), y + I(0.42), rw - I(0.4), h - I(0.5), body, 13, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        y += h + I(0.15)

    takeaway(s, I(5.82), I(0.9), "TAKEAWAY",
             "When C_{2} and C_{3} learn the errors of C_{1}′, both errors decrease. This result supports the first "
             "cause. Retraining all downstream models after every update has a cost. This cost motivates RQ3.",
             14.5)

    notes(s, 'If the first cause is right, retraining C2 and C3 on the output of C1 prime should help. We test '
             'this idea on one retraining of each original C1. The red bars show the errors after the retraining '
             'with C2 and C3 fixed. The navy bars show the errors after we retrain C2 and C3. The prediction error'
             ' decreases by 0.36 and 0.16 metres. The planning error decreases by 0.44 and 0.30 metres. Both '
             'errors fall below their values before the retraining. This result supports the first cause. However,'
             ' retraining every downstream model after every update has a cost. This cost motivates RQ3.')


def result_detection(prs):
    s = content_slide(prs, "Result: Detection Without Ground Truth", "IV · Results · RQ3")
    I = Inches

    s.shapes.add_picture(str(HERE / "fig_detection.png"), int(MARGIN - I(0.1)), int(I(1.35)), width=int(I(7.1)))

    rx = MARGIN + I(7.2)
    rw = W - MARGIN - rx
    cards = [("WHY",
              "After deployment, new data have no ground truth. A warning sign must work without ground truth."),
             ("WHAT WE MEASURE",
              "Nine measures compare the outputs before and after the retraining on the same test images. "
              "Seven use the output of C_{1}. Two use the output of C_{2}."),
             ("WHY RANK CORRELATION",
              "A useful measure is larger for retrainings with a larger Δ_{2}. Rank correlation measures this "
              "agreement. A value of +1 is perfect agreement, and 0 is none."),
             ("RESULT",
              "The best measure is the shift of predicted positions. Its rank correlation is **+0.68** over all "
              "30 retrainings, but only **+0.29** and **−0.28** within each original C_{1}.")]
    hs = [I(0.82), I(1.0), I(1.08), I(1.08)]
    y = I(1.4)
    for (head, body), h in zip(cards, hs):
        shape(s, MSO_SHAPE.RECTANGLE, rx, y, rw, h, fill=IVORY)
        bar(s, rx, y, Pt(4), h, GOLD)
        text(s, rx + I(0.25), y + I(0.08), rw - I(0.4), I(0.28), head, 10.5, GOLD, bold=True,
             align=PP_ALIGN.LEFT, spacing=1.5)
        text(s, rx + I(0.25), y + I(0.36), rw - I(0.4), h - I(0.4), body, 12.5, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
        y += h + I(0.08)

    takeaway(s, I(5.95), I(0.82), "ANSWER TO RQ3",
             "H3 is not supported. The best measure only separates the two original C_{1}. "
             "No measure identifies single retrainings with entangled enhancement.", 14.5)

    notes(s, 'For RQ3, we test nine measures that compare the outputs before and after the retraining on the same '
             'images. Seven use the output of C1, such as the number of objects per image. Two use the output of '
             'C2, such as the shift of the predicted positions. A useful measure is larger when delta two, the '
             'increase of the prediction error, is larger. Rank correlation measures this agreement. A value of '
             'one is perfect, and zero is none. The best measure, the shift of the predicted positions, reaches '
             '0.68 over all retrainings, but only 0.29 and minus 0.28 within each original C1. The measure '
             'separates the two original C1, not single retrainings. H3 is therefore not supported.')


def conclusion(prs):
    s = content_slide(prs, "Conclusion", "V · Conclusion")
    I = Inches
    tw = W - 2 * MARGIN

    # what we did
    sy, sh_ = I(1.45), I(1.05)
    shape(s, MSO_SHAPE.RECTANGLE, MARGIN, sy, tw, sh_, fill=IVORY)
    bar(s, MARGIN, sy, Pt(6), sh_, GOLD)
    text(s, MARGIN + I(0.35), sy, tw - I(0.6), sh_,
         "In this study, we empirically investigated entangled enhancement in an autonomous driving pipeline of "
         "three ML models. We retrained the perception model C_{1} 30 times and trained C_{2} and C_{3} in two "
         "training modes.", 16, INK, align=PP_ALIGN.LEFT)

    rows = [("RQ1 · OCCURRENCE",
             "We observe that a retraining that improves C_{1} increases the prediction error of C_{2}. The "
             "increase is above seed variation in 11 of 30 retrainings and significant in 3. The planning error "
             "of C_{3} increases by at most 0.13 m.",
             "H1 SUPPORTED\nFOR C_{2}", GOLD),
            ("RQ2 · CAUSE",
             "We confirm that the increase in C_{2} comes from the adaptation of C_{2} to the errors of C_{1}. "
             "The increase appears only after C_{1}-output training. Retraining C_{2} and C_{3} on the output of "
             "C_{1}′ lowers both errors.",
             "H2 SUPPORTED\nFOR C_{2}", GOLD),
            ("RQ3 · DETECTION",
             "We find no measure without ground truth that identifies single retrainings with entangled "
             "enhancement.",
             "H3 NOT\nSUPPORTED", WINE)]
    y, rh, gap = I(2.7), I(1.15), I(0.13)
    lab_w, ver_w = I(2.3), I(1.9)
    for label, body, verdict, vcol in rows:
        bar(s, MARGIN, y, lab_w, rh, NAVY)
        text(s, MARGIN + I(0.25), y, lab_w - I(0.3), rh, label, 12.5, GOLD, bold=True, align=PP_ALIGN.LEFT,
             spacing=1.2)
        shape(s, MSO_SHAPE.RECTANGLE, MARGIN + lab_w, y, tw - lab_w, rh, fill=WHITE, line=LIGHT_GOLD, line_w=0.75)
        text(s, MARGIN + lab_w + I(0.3), y, tw - lab_w - ver_w - I(0.6), rh, body, 13.5, INK, align=PP_ALIGN.LEFT)
        vx = W - MARGIN - ver_w - I(0.15)
        bar(s, vx, y + I(0.2), ver_w, rh - I(0.4), vcol)
        text(s, vx, y + I(0.2), ver_w, rh - I(0.4), verdict, 11.5, WHITE, bold=True, spacing=1)
        y += rh + gap

    notes(s, 'In this study, we investigated entangled enhancement in an autonomous driving pipeline of three ML '
             'models. For RQ1, we observe that a retraining that improves C1 increases the prediction error of C2,'
             ' clearly in 11 of 30 retrainings. The planning error increases by at most 0.13 metres. For RQ2, we '
             'confirm that the increase in C2 comes from the adaptation of C2 to the errors of C1. For RQ3, we '
             'find no measure without ground truth that identifies single retrainings with entangled enhancement.')


def future_work(prs):
    s = content_slide(prs, "Future Work", "V · Conclusion")
    I = Inches
    tw = W - 2 * MARGIN

    # main next step
    my, mh = I(1.5), I(2.0)
    bar(s, MARGIN, my, tw, mh, NAVY)
    bar(s, MARGIN, my, Pt(6), mh, GOLD)
    text(s, MARGIN + I(0.4), my + I(0.2), I(4), I(0.35), "NEXT STEP", 12.5, GOLD, bold=True, align=PP_ALIGN.LEFT,
         spacing=2)
    text(s, MARGIN + I(0.4), my + I(0.55), tw - I(0.8), I(0.5), "Early detection of entangled enhancement", 24,
         WHITE, bold=True, align=PP_ALIGN.LEFT)
    text(s, MARGIN + I(0.4), my + I(1.1), tw - I(0.8), I(0.8),
         "We are developing a method that detects entangled enhancement before C_{1}′ is deployed, "
         "and before ground truth for the new data exists. None of the nine measures in this study achieves this "
         "detection.",
         15.5, PALE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

    cards = [("More test scenes",
              "With 50 test scenes, only 3 of 30 increases are statistically significant. "
              "More scenes give narrower intervals."),
             ("More original C_{1}",
              "All 11 increases above seed variation come from one original C_{1}. "
              "More original C_{1} show how often this case occurs."),
             ("Input to retraining models",
              "The measured frequency of entangled enhancement can serve as the probability in the models "
              "of Wang and Machida.")]
    gap = I(0.3)
    cw = (tw - 2 * gap) / 3
    cy, ch = I(3.85), I(2.2)
    for k, (head, body) in enumerate(cards):
        cx = MARGIN + k * (cw + gap)
        shape(s, MSO_SHAPE.RECTANGLE, cx, cy, cw, ch, fill=IVORY)
        bar(s, cx, cy, cw, Pt(4), GOLD)
        num_badge(s, cx + I(0.25), cy + I(0.3), I(0.5), str(k + 1), size=15)
        text(s, cx + I(0.9), cy + I(0.3), cw - I(1.1), I(0.5), head, 17, NAVY, bold=True, align=PP_ALIGN.LEFT)
        text(s, cx + I(0.25), cy + I(1.0), cw - I(0.5), ch - I(1.1), body, 14, INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)

    notes(s, 'Our next step is the early detection of entangled enhancement, before C1 prime is deployed and '
             'before ground truth exists. We also plan more test scenes, because only 3 of the 30 increases are '
             'significant with 50 scenes. We plan more original C1, because all 11 clear increases come from one '
             'original C1. Finally, the measured frequency can serve as the probability in the models of Wang and '
             'Machida.')


def thanks(prs):
    global TITLE
    TITLE = "Thank You"
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s, 0, 0, W, H, NAVY)
    inset = Inches(0.35)
    shape(s, MSO_SHAPE.RECTANGLE, inset, inset, W - 2 * inset, H - 2 * inset, line=GOLD, line_w=1)
    shape(s, MSO_SHAPE.RECTANGLE, inset + Inches(0.08), inset + Inches(0.08),
          W - 2 * inset - Inches(0.16), H - 2 * inset - Inches(0.16), line=GOLD, line_w=0.5)
    text(s, Inches(1), Inches(2.3), W - Inches(2), Inches(1.0), "Thank You", 48, WHITE, bold=True)
    ornament(s, W / 2, Inches(3.55), Inches(2.0))
    text(s, Inches(1), Inches(4.0), W - Inches(2), Inches(0.5), "Questions and comments are welcome.", 20, PALE,
         italic=True)
    text(s, Inches(1), Inches(5.4), W - Inches(2), Inches(0.4), "Rashedul Arefin Ifty  ·  Dependable AI Systems Laboratory  ·  University of Tsukuba",
         15, GOLD, spacing=1)
    notes(s, "Thank you for your attention. I welcome your questions and comments.")


PINK = RGBColor(0xF3, 0xE1, 0xE1)  # cell shading: increase larger than seed variation


def fmt(v):
    if round(v, 3) == 0:
        return "0.000"
    return f"{v:+.3f}".replace("-", "−")


def appendix_retrainings(X, letter):
    """All 15 retrainings of the original C1 with seed X, numbered as in the result figures."""
    def make(prs):
        import make_figures as mf
        s = content_slide(prs, f"Appendix {letter}: Retrainings of Original C_{{1}}, Seed {X}", "Appendix")
        I = Inches
        rows = [["#", "DATA", "EPOCHS", "SEED", "δ_{1}", "Δ_{2}  C_{1}-OUTPUT", "Δ_{2}  GT", "Δ_{3}  C_{1}-OUTPUT",
                 "Δ_{3}  GT"]]
        shade = []
        for k, r in enumerate(mf.rows, 1):
            if r["base"] != X:
                continue
            cells = [str(k), f"{int(r['frac'] * 100)} %", str(r["epochs"]), str(r["seed"]), fmt(r["d1"])]
            for key, sig, tau in (("d2c", "s2c", mf.TAU2), ("d2g", "s2g", mf.TAU2),
                                  ("d3c", "s3c", mf.TAU3), ("d3g", "s3g", mf.TAU3)):
                v = r[key]
                cells.append(("**" + fmt(v) + " *" + "**") if r[sig] else fmt(v))
                if v > tau:
                    shade.append((len(rows), len(cells) - 1))
            rows.append(cells)
        tw = W - 2 * MARGIN
        cw = [I(0.6), I(0.95), I(0.95), I(0.8), I(1.15)]
        rest = (tw - sum(cw)) / 4
        tbl = table(s, MARGIN, I(1.4), cw + [rest] * 4, rows, I(0.3), I(0.42), size=12)
        for i, j in shade:
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = PINK
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = WINE
        text(s, MARGIN, I(6.37), tw, I(0.3),
             f"Changes in metres after the retraining.  Shaded: larger than seed variation "
             f"({mf.TAU2:.3f} m for Δ_{{2}}, {mf.TAU3:.3f} m for Δ_{{3}}).  * Statistically significant.  "
             f"Data: share of Singapore 1 used for the retraining.", 11, GREY, italic=True, align=PP_ALIGN.LEFT)
        notes(s, f"This appendix slide lists the 15 retrainings of the original C1 with seed {X}. "
                 "The numbers match the result figures.")
    make.__name__ = f"appendix_{letter}"
    return make


def appendix_models(prs):
    s = content_slide(prs, "Appendix C: Details of the ML Models", "Appendix")
    I = Inches
    tw = W - 2 * MARGIN
    cards = [("C_{1} · Perception model",
              ["YOLOv11n object detector",
               "Input: front-camera image, 640 pixels",
               "Trained 10 epochs on Boston 1, batch size 8",
               "Last checkpoint used; the test scenes never select a model",
               "mAP50 on the test scenes: 0.109 (seed 0), 0.117 (seed 1)"]),
             ("C_{2} · Prediction model",
              ["AutoBot-Ego style transformer, centred on each target object",
               "Input: detections in the last four labelled images, relative to the target",
               "2 encoder and 2 decoder layers, 128 features, 4 attention heads",
               "Output: six possible paths of 3 s",
               "Same object across images found by attention, without a separate matching step",
               "736,799 parameters"]),
             ("C_{3} · Planning model",
              ["PlanT style transformer, one input token per object",
               "Object token: position, most likely path from C_{2}, class, confidence",
               "Further tokens: driving command and ego path of the last 1.5 s",
               "4 layers, 128 features, 4 attention heads",
               "Recurrent decoder: six points of the planned path, 0.5 s apart",
               "636,930 parameters"])]
    gap = I(0.3)
    cw = (tw - 2 * gap) / 3
    cy, ch = I(1.45), I(3.95)
    for k, (head, items) in enumerate(cards):
        cx = MARGIN + k * (cw + gap)
        shape(s, MSO_SHAPE.RECTANGLE, cx, cy, cw, ch, fill=WHITE, line=NAVY, line_w=1)
        bar(s, cx, cy, cw, I(0.5), NAVY)
        text(s, cx + I(0.2), cy, cw - I(0.4), I(0.5), head, 15, WHITE, bold=True, align=PP_ALIGN.LEFT)
        text(s, cx + I(0.2), cy + I(0.65), cw - I(0.35), ch - I(0.75), "\n".join("•  " + t for t in items), 13.5,
             INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15)
    ty, th = I(5.6), I(1.0)
    shape(s, MSO_SHAPE.RECTANGLE, MARGIN, ty, tw, th, fill=IVORY)
    bar(s, MARGIN, ty, Pt(6), th, GOLD)
    text(s, MARGIN + I(0.35), ty, I(2.0), th, "TRAINING OF\nC_{2} AND C_{3}", 12, GOLD, bold=True,
         align=PP_ALIGN.LEFT, spacing=1.2)
    text(s, MARGIN + I(2.4), ty, tw - I(2.6), th,
         "Boston 2: 86 training and 14 validation scenes. Early stopping on the validation scenes (patience 8 "
         "epochs for C_{2}, 10 for C_{3}). Mirror and rotation (±10°) augmentation. The error of C_{2} uses the "
         "objects of C_{1} matched one to one with a ground-truth object of the same class.",
         13, INK, align=PP_ALIGN.LEFT)
    notes(s, "This appendix slide gives the details of the three ML models and the training of C2 and C3.")


SLIDES = [cover, outline, background, existing_studies, objective, pipeline, training_modes, setup, evaluation,
          model_check, result_prediction, result_planning, result_retraining, result_detection, conclusion,
          future_work, thanks, appendix_retrainings(0, "A"), appendix_retrainings(1, "B"), appendix_models]


def main(only=None, out=OUT):
    """Build all slides, or with only=N a one-slide deck (for previews) that keeps slide number N."""
    global NUM
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for NUM, make in enumerate(SLIDES, 1):
        if only is None or NUM == only:
            make(prs)
    prs.save(out)
    print("wrote", out, f"({len(prs.slides)} slides)")
    if only is None:
        write_script()


def plain(t):
    """Slide-title markup to readable Markdown: C_{1} -> C₁."""
    sub = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")
    return re.sub(r"_\{([^}]*)\}", lambda m: m.group(1).translate(sub), t).replace("**", "")


QUESTIONS = [  # expected audience questions, with answers taken from the results and the code
    ("What explains the increase of the planning error after GT training?",
     "The first cause cannot explain it. A GT-trained C3 learns from the true objects and their true future paths. "
     "At test time, the same C3 receives the detected objects and the paths that C2 predicts. A change in the "
     "objects that C1 prime finds can therefore move the plan of this C3 more. This study does not test this "
     "explanation directly, and the increase is at most 0.13 metres."),
    ("Why do all clear increases come from the original C1 with seed 0?",
     "The C2 trained on the output of this C1 is sensitive to a change in the errors of C1. When we replace this C1 "
     "by the other original C1, which is almost equally accurate (0.009 apart in mAP50), the error of this C2 "
     "increases by 0.112 metres. The C2 trained on the output of the seed-1 C1 improves by 0.098 metres under the "
     "same replacement. The adaptation of C2 therefore differs between the two original C1."),
    ("Two measures on slide 14 correlate well within seed 0. Why is H3 still not supported?",
     "Within seed 0, the number of objects per image reaches plus 0.58, and the largest gap of the confidence values "
     "reaches minus 0.56. Within seed 1, the same measures reach only minus 0.14 and minus 0.09. A useful warning "
     "sign must work for every original C1, and 15 retrainings per original C1 are a small sample."),
    ("Does the ego path let C3 ignore the objects?",
     "An ego-path input can let a planning model ignore the objects on nuScenes, which Li et al. and Zhai et al. "
     "report. Without the ego path, C3 is not better than the average Boston path. With the ego path, removing the "
     "objects still moves the plan by 0.73 to 1.65 metres, which is above our threshold of 0.5 metres."),
    ("Why is the accuracy of C1 low, with an mAP50 of about 0.11?",
     "C1 is trained on only 100 Boston scenes with the front camera, and the test scenes come from another city. "
     "This study needs the change of accuracy, not a high accuracy. Every one of the 30 retrainings improves mAP50 "
     "by 0.048 to 0.118."),
    ("Why only 350 scenes and 50 test scenes?",
     "We reduced the data to 350 scenes to run 30 retrainings with two training modes on one GPU. "
     "With 50 test scenes, only 3 of the 30 increases are statistically significant. More test scenes are part "
     "of the future work."),
    ("Does the accuracy gain delta one predict the increase?",
     "No. The rank correlation between delta one and delta two is plus 0.04. A larger accuracy gain does not "
     "come with a larger increase of the prediction error."),
]


def write_script(wpm=130):
    main_part = [x for x in SCRIPT if not x[1].startswith("Appendix")]
    words = sum(len(t.split()) for _, _, t in main_part)
    L = ["# Speaker Script: An Empirical Analysis of Entangled Enhancement in an Autonomous Driving Pipeline", "",
         f"Rashedul Arefin Ifty · 0ALD515 Research in Computer Science B · University of Tsukuba", "",
         f"Main talk: slides 1 to {len(main_part)}, {words} words, about {words / 110:.1f} minutes at a slow 110 "
         f"words per minute and {words / wpm:.1f} minutes at {wpm}. The appendix slides are shown only when a "
         f"question needs them.", "",
         "Generated by build_slides.py together with CS_presentation_IFTY.pptx. Edit the script in "
         "build_slides.py and rebuild. The slides and the script then stay in sync.", ""]
    appendix_started = False
    for n, title, t in SCRIPT:
        if title.startswith("Appendix") and not appendix_started:
            L += ["---", "", "# Appendix (only if asked)", ""]
            appendix_started = True
        sec = len(t.split()) / wpm * 60
        L += [f"## Slide {n}: {plain(title)}  (about {sec:.0f} s)", "", t, ""]
    L += ["---", "", "# Expected Questions", "",
          "Short answers for the discussion. Every number comes from the results shown in the slides or the "
          "appendix.", ""]
    for k, (q, a) in enumerate(QUESTIONS, 1):
        L += [f"**Q{k}. {q}**", "", a, ""]
    SCRIPT_OUT.write_text("\n".join(L))
    print("wrote", SCRIPT_OUT)


if __name__ == "__main__":
    import sys
    if len(sys.argv) == 3:  # build_slides.py N preview.pptx
        main(int(sys.argv[1]), Path(sys.argv[2]))
    else:
        main()
